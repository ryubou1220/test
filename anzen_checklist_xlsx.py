#!/usr/bin/env python3
"""橋梁新設工事 安全対策チェックリスト（案）を Excel に書き出す。"""
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "out/anzen_checklist.xlsx"
FONT = "游ゴシック"

SECTIONS = [
    ("1. 共通・安全管理体制", [
        "施工計画書・作業手順書があり、当日の作業内容と合っている",
        "元請・協力会社の安全管理体制表と緊急連絡網が現場に掲示されている",
        "新規入場者教育を実施し、記録を残している",
        "朝礼・TBM・KY活動で当日の危険と対策を全員で確認している",
        "必要な作業主任者を選任し、氏名と職務を掲示している（足場、地山掘削、土止め支保工、型枠支保工、鋼橋・コンクリート橋架設など）",
        "資格が必要な作業は、有資格者が従事し、資格証の写しを揃えている",
        "保護具（ヘルメット、フルハーネス、安全靴、保護メガネなど）を正しく着用している",
        "作業員の体調を朝に確認している（睡眠不足、飲酒、体調不良）",
        "高齢者・外国人労働者・未熟練者に配慮した配置と指示をしている",
        "救急箱・AEDの置き場所と、最寄りの病院への連絡方法を全員が知っている",
        "現場内の整理・整頓がされ、通路が確保されている",
        "元請と協力会社が作業間の連絡調整（上下作業・同時作業の調整）をしている",
    ]),
    ("2. 仮設・足場・墜落防止", [
        "高さ 2m 以上の作業箇所に作業床がある（幅 40cm 以上、床材のすき間 3cm 以下）",
        "足場に手すり（高さ 85cm 以上）と中さん（35〜50cm）がある",
        "幅木（高さ 10cm 以上）やメッシュシートで物の落下を防いでいる",
        "足場の壁つなぎ・控え・筋かいに外れや緩みがない",
        "足場の作業開始前点検を、指名された点検者が行い、記録している",
        "強風・大雨・大雪・中震（震度4）以上の地震の後や、組立・変更の後に足場を点検している",
        "手すりのない箇所では、親綱を張り、フルハーネス型墜落制止用器具を使っている",
        "フルハーネス使用者が特別教育を受けている",
        "開口部・桁端・橋台天端に囲い・覆い・手すりがある",
        "昇降設備（階段・昇降設備・はしご）があり、固定されている",
        "吊り足場・張出し足場の吊り材や支持金具に損傷や外れがない",
        "足場の最大積載荷重を表示し、資材を置きすぎていない",
        "高所作業車は作業開始前点検をし、乗員が墜落制止用器具を使っている",
        "足場の組立・解体中は立入禁止とし、作業主任者が指揮している",
    ]),
    ("3. 下部工（基礎・掘削・土留め・橋台・橋脚）", [
        "掘削の前に、地山の状態・湧水・埋設物を調査している",
        "掘削の法勾配が、地山の種類と掘削の高さに合っている",
        "掘削の天端に資材や残土を置かず、重機は法肩から離れている",
        "法面に亀裂・湧水・浮石がない（作業開始前と降雨後に確認）",
        "土止め支保工（切梁・腹起し）に変形・緩み・脱落がない（設置後 7日以内ごと、中震以上・大雨の後に点検）",
        "掘削箇所の周りに転落防止柵と昇降設備がある",
        "杭打機・杭抜機の据付地盤が平坦で、敷鉄板などで転倒を防いでいる",
        "杭打機の組立て・解体では、作業指揮者を置いている",
        "杭打ち中は、杭やリーダーの周りを立入禁止にしている",
        "型枠支保工の組立図があり、そのとおり組まれている",
        "コンクリート打設中に、型枠・支保工の変形を監視する人を配置している",
        "コンクリートポンプ車のブーム下を立入禁止にし、アウトリガーを最大に張り出している",
        "鉄筋の端部にキャップを付け、刺さりを防いでいる",
        "深礎・ケーソンなどの内部に入る前に、酸素濃度（18%以上）と有害ガスを測定している",
    ]),
    ("4. 上部工（架設）", [
        "架設計画書（架設ステップ、吊り荷重、作業半径）があり、全員に周知している",
        "鋼橋架設等またはコンクリート橋架設等の作業主任者が、現場で直接指揮している",
        "移動式クレーンの定格荷重を、作業半径ごとに確認している（吊荷重量＋吊具の重量）",
        "クレーンの据付地盤の耐力を確認し、敷鉄板を敭いている。アウトリガーを最大に張り出している",
        "クレーンの作業開始前点検（過負荷防止装置、巻過防止装置、ブレーキ）をしている",
        "クレーン運転士と玉掛け者が資格を持ち、合図者を一人に決めている",
        "玉掛け用ワイヤー・シャックル・吊クランプに損傷がなく、安全係数を満たしている",
        "吊り荷の下と旋回範囲を立入禁止にし、介錯ロープで荷を誘導している",
        "ベント（仮支柱）の基礎と部材に沈下・傾き・変形がない",
        "送出し架設では、反力・たわみ・速度を計測し、管理値を超えたら止める手順がある",
        "架設桁・トラベラークレーンなどの特殊設備に、逸走・転倒の防止装置がある",
        "架設した桁に、一時的な転倒防止（仮固定・横構の仮締め）をしている",
        "桁上では、親綱とフルハーネスを使い、移動中も常にどこかに掛けている",
        "床版の型枠・支保工や、張出し床版の端部に手すりがある",
        "PC緊張作業中は、ジャッキの後方を立入禁止にし、防護板を置いている",
        "ボルト締め・溶接などで使う工具に落下防止のヒモを付けている",
    ]),
    ("5. 河川・水上作業", [
        "出水時の退避基準（水位、雨量、警報・注意報）を数字で決め、掲示している",
        "上流の水位観測所・雨量を、作業中に定期的に確認している",
        "退避経路と退避場所を決め、全員が知っている",
        "河川内に置いた重機・資材を、出水の前に撤去する手順がある",
        "仮締切・仮橋に洗掘・変形・漏水がない",
        "水辺・水上で作業する人が救命胴衣を着ている",
        "救命浮環・ロープ・救命ボートを用意し、監視人を置いている",
        "作業船は定員・積載量を守り、係留が確実にされている",
        "河川管理者の占用許可条件（作業期間・出水期の制限）を守っている",
        "濁水・油の流出を防ぐ対策（汚濁防止膜、オイルフェンス）をしている",
    ]),
    ("6. 重機・車両・電気・火気", [
        "車両系建設機械の作業計画（機種、運行経路、作業方法）を作り、周知している",
        "運転者が資格を持ち、作業開始前点検をしている",
        "重機の作業半径内を立入禁止にし、誘導員を配置している",
        "後退時の警報装置・接近警報装置が動いている",
        "バックホウで荷を吊るときは、クレーン機能付きの機種を使い、用途外使用をしていない",
        "路肩・法肩での転落を防ぐため、路肩の表示や車止めがある",
        "架空線の位置を確認し、離隔距離の確保、防護管の設置、監視人の配置をしている",
        "地下埋設物（ガス・水道・通信・電力）を管理者と確認し、試掘している",
        "仮設電気設備に漏電遮断器とアースがある",
        "ケーブルに損傷がなく、通路を横切る箇所は防護している",
        "交流アーク溶接機に自動電撃防止装置があり、動作を確認している",
        "溶接・ガス切断の周りに燃えやすい物がなく、消火器と火花よけを置いている",
        "ガスボンベを立てて固定し、直射日光を避けている",
        "閉じた場所（箱桁の中など）で作業するときは、換気と酸素・有害ガスの測定をしている",
        "エンジン発電機などを屋内や箱桁の中で使っていない（一酸化炭素中毒の防止）",
    ]),
    ("7. 第三者災害防止・交通誘導", [
        "現場の周りを仮囲いで囲い、部外者が入れないようにしている",
        "道路使用許可・占用許可の条件（時間帯、規制方法）を守っている",
        "交通規制の標識・バリケード・保安灯を、許可された規制図どおりに置いている",
        "交通誘導警備員を必要な人数だけ配置している（検定合格警備員が必要な路線かを確認）",
        "歩行者・自転車の通路を確保し、安全に誘導している",
        "既設道路の上や近くで吊り作業をするときは、一時的に通行を止めている",
        "桁下や道路の上に、落下物防護のための防護棚・防護ネットがある",
        "工事車両が出入口で一旦停止し、誘導員が案内している",
        "工事用道路の無理な走行（進入時の速度超過、過積載）をしていない",
        "工事車両のタイヤ洗浄や散水で、泥や粉じんを外に出していない",
        "作業終了後に、仮囲いの施錠、重機のキー管理、開口部の閉鎖をしている",
        "近隣へ作業内容・時間を事前に知らせ、騒音・振動を抑えている",
    ]),
    ("8. 気象・自然災害・熱中症", [
        "作業中止基準（「作業中止基準」シート）を現場で決め、掲示している",
        "風速計があり、風速を記録している",
        "気象情報（警報・注意報、雨雲レーダー）を朝と作業中に確認している",
        "強風の前に、シート類のたたみ込みや、資材・型枠の飛散防止をしている",
        "WBGT を測定し、作業員に知らせている",
        "日陰の休憩所・飲み物・塩分・冷却用品を用意している",
        "熱中症の疑いがある人を見つけたときの報告先と対応手順（作業を外す、冷やす、医療機関へ運ぶ）を決め、周知している",
        "暑さに慣れていない人（新規入場者、休み明け）の作業量を減らしている",
        "地震・津波のときの退避場所と連絡方法を決めている",
    ]),
]
# 4-4 の誤字を修正（敭→敷）
SECTIONS[3][1][3] = SECTIONS[3][1][3].replace("敭", "敷")

STOP_CRITERIA = [
    ("強風", "10分間の平均風速 10m/s 以上", "高所作業・クレーン作業を中止。クレーンのジブを下げる"),
    ("大雨", "1回の降雨量 50mm 以上", "高所作業を中止。その後、足場・法面・土止めを点検"),
    ("大雪", "1回の降雪量 25cm 以上", "高所作業を中止。その後、足場を点検"),
    ("地震", "中震（震度4）以上", "作業を止めて退避。その後、足場・支保工・ベントを点検"),
    ("雷", "雷鳴が聞こえる、または雷注意報", "高所・金属を扱う作業を中止し、建物や車内へ退避"),
    ("出水", "現場で決めた水位・雨量（5章）", "河川内の作業を中止し、重機・資材を撤去"),
    ("暑さ", "WBGT 28 以上、または気温 31℃ 以上", "休憩を増やす。熱中症の報告体制と対応手順を整備（8-7）"),
]

LAWS = [
    ("労働安全衛生法", "1章（安全管理体制、作業主任者、教育）"),
    ("労働安全衛生規則", "2・3・4・6・8章（墜落防止、足場、掘削、型枠支保工、橋梁架設、車両系建設機械、電気、熱中症）"),
    ("クレーン等安全規則", "4章（移動式クレーン、玉掛け、強風時の作業中止）"),
    ("酸素欠乏症等防止規則", "3・6章（ケーソン・箱桁などの閉じた場所）"),
    ("道路交通法・道路法、警備業法", "7章（道路使用・占用、交通誘導警備）"),
    ("河川法", "5章（河川占用の条件）"),
    ("建設工事公衆災害防止対策要綱（土木工事編）", "7章（第三者災害の防止）"),
    ("土木工事安全施工技術指針", "全般"),
]

# ---- スタイル ---------------------------------------------------------------
thin = Side(style="thin", color="808080")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
EX_FILL = PatternFill("solid", fgColor="F2F2F2")
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def header_row(ws, row, labels):
    for c, lab in enumerate(labels, 1):
        cell = ws.cell(row, c, lab)
        cell.font = f(True, 10, "FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def page_setup(ws, title_rows=None, landscape=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.center.text = "&P / &N"


wb = Workbook()

# ---- チェックリスト ------------------------------------------------------------
ws = wb.active
ws.title = "チェックリスト"
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 58
ws.column_dimensions["C"].width = 7
ws.column_dimensions["D"].width = 30

ws["A1"] = "橋梁新設工事 安全対策チェックリスト（案）"
ws["A1"].font = f(True, 14)
ws.merge_cells("A1:D1")

info = ["工事名", "施工箇所", "点検日", "点検者（職・氏名）", "立会者", "当日の主な作業"]
r = 3
for lab in info:
    ws.cell(r, 1, lab).font = f(True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=1)
    ws.cell(r, 2).fill = INPUT_FILL
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c in range(1, 5):
        ws.cell(r, c).border = BORDER
    ws.cell(r, 2).font = f()
    r += 1
ws["B5"].number_format = "yyyy/mm/dd"

# 集計（判定列を数える）
r += 1
summary_row = r
ws.cell(r, 1, "集計").font = f(True)
labels = [("○ 適正", "○"), ("× 不適", "×"), ("－ 該当なし", "－")]
r += 1
first_item_row_placeholder = r  # 後で範囲を入れる
summary_cells = []
for lab, sym in labels:
    ws.cell(r, 1, sym).alignment = CENTER
    ws.cell(r, 2, lab)
    summary_cells.append((r, sym))
    r += 1
ws.cell(r, 1, "未").alignment = CENTER
ws.cell(r, 2, "未記入")
blank_row = r
r += 1
ws.cell(r, 1, "計").alignment = CENTER
ws.cell(r, 2, "全項目数")
total_row = r
r += 2

ws.cell(r, 1, "判定欄は ○・×・－ をリストから選択。×の項目は「是正記録」シートに転記する。黄色のセルが記入欄。").font = f(size=9, color="595959")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
r += 2

head = r
header_row(ws, head, ["No.", "点検項目", "判定", "是正内容・備考"])
r += 1
first = r
for si, (title, items) in enumerate(SECTIONS, 1):
    ws.cell(r, 1, title).font = f(True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    for c in range(1, 5):
        ws.cell(r, c).fill = SEC_FILL
        ws.cell(r, c).border = BORDER
    r += 1
    for ii, text in enumerate(items, 1):
        vals = [f"{si}-{ii}", text, None, None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = f()
            cell.border = BORDER
            cell.alignment = CENTER if c in (1, 3) else WRAP
        ws.cell(r, 3).fill = INPUT_FILL
        ws.cell(r, 4).fill = INPUT_FILL
        ws.row_dimensions[r].height = 16 * -(-len(text) // 30) + 4
        r += 1
last = r - 1

rng = f"$C${first}:$C${last}"
for row, sym in summary_cells:
    ws.cell(row, 3, f'=COUNTIF({rng},"{sym}")')
ws.cell(total_row, 3, f'=COUNTA($B${first}:$B${last})')
ws.cell(blank_row, 3, f"=C{total_row}-SUM(C{summary_cells[0][0]}:C{summary_cells[-1][0]})")
for row in range(summary_cells[0][0], total_row + 1):
    for c in range(1, 4):
        ws.cell(row, c).border = BORDER
        ws.cell(row, c).font = f(bold=(row == total_row))
    ws.cell(row, 3).alignment = CENTER

dv = DataValidation(type="list", formula1='"○,×,－"', allow_blank=True,
                    error="○・×・－ から選んでください", errorTitle="判定")
ws.add_data_validation(dv)
dv.add(rng.replace("$", ""))
red = PatternFill("solid", fgColor="F8CBAD")
ws.conditional_formatting.add(rng.replace("$", ""),
                              CellIsRule(operator="equal", formula=['"×"'], fill=red,
                                         font=Font(name=FONT, bold=True, color="C00000")))
ws.conditional_formatting.add(f"C{summary_cells[1][0]}",
                              FormulaRule(formula=[f"C{summary_cells[1][0]}>0"], fill=red))
ws.freeze_panes = ws.cell(head + 1, 1)
page_setup(ws, f"{head}:{head}")

# ---- 作業中止基準 --------------------------------------------------------------
ws2 = wb.create_sheet("作業中止基準")
ws2["A1"] = "作業中止基準（目安）"
ws2["A1"].font = f(True, 14)
ws2["A2"] = "一般的な目安。橋の形式や架設計画に合わせて現場ごとに決め、朝礼で確認する。"
ws2["A2"].font = f(size=9, color="595959")
header_row(ws2, 4, ["気象の状態", "基準の目安", "現場で決めた基準", "対応"])
for i, (a, b, d) in enumerate(STOP_CRITERIA, 5):
    for c, v in enumerate([a, b, None, d], 1):
        cell = ws2.cell(i, c, v)
        cell.font = f()
        cell.border = BORDER
        cell.alignment = CENTER if c == 1 else WRAP
    ws2.cell(i, 3).fill = INPUT_FILL
    ws2.row_dimensions[i].height = 32
for col, w in zip("ABCD", (12, 30, 22, 46)):
    ws2.column_dimensions[col].width = w
page_setup(ws2, landscape=True)

# ---- 是正記録 ----------------------------------------------------------------
ws3 = wb.create_sheet("是正記録")
ws3["A1"] = "点検結果・是正記録"
ws3["A1"].font = f(True, 14)
ws3["A2"] = "「チェックリスト」で × が付いた項目を転記し、是正が終わるまで追う。灰色の行は記入例（使うときは消す）。"
ws3["A2"].font = f(size=9, color="595959")
cols = ["項目No.", "不適の内容", "是正の内容", "担当者", "期限", "是正日", "確認者", "状況"]
header_row(ws3, 4, cols)
example = ["2-2", "P3橋脚足場の北面で中さんが1スパン外れていた", "中さんを再設置し、同じ面の全スパンを再点検",
           "鈴木（○○組）", "2026/10/09", "2026/10/09", "田中"]
for c, v in enumerate(example, 1):
    cell = ws3.cell(5, c, v)
    cell.font = f(italic=True, color="595959")
    cell.fill = EX_FILL
    cell.border = BORDER
    cell.alignment = WRAP
ws3.row_dimensions[5].height = 32
for rr in range(6, 26):
    for c in range(1, 9):
        cell = ws3.cell(rr, c)
        cell.border = BORDER
        cell.font = f()
        cell.alignment = WRAP
        if c < 8:
            cell.fill = INPUT_FILL
    ws3.row_dimensions[rr].height = 28
for rr in range(5, 26):
    ws3.cell(rr, 8, f'=IF(A{rr}="","",IF(F{rr}<>"","完了",IF(AND(ISNUMBER(E{rr}),E{rr}<TODAY()),"期限超過","対応中")))')
    ws3.cell(rr, 8).alignment = CENTER
    for c in (5, 6):
        ws3.cell(rr, c).number_format = "yyyy/mm/dd"
ws3.conditional_formatting.add("H5:H25", CellIsRule(operator="equal", formula=['"期限超過"'], fill=red))
for col, w in zip("ABCDEFGH", (9, 34, 34, 14, 12, 12, 10, 10)):
    ws3.column_dimensions[col].width = w

ws3["A28"] = "確認欄"
ws3["A28"].font = f(True, 11)
header_row(ws3, 29, ["役割", "氏名", "確認日"])
for i, role in enumerate(["点検者", "現場代理人", "安全担当者", "監理技術者"], 30):
    ws3.cell(i, 1, role).font = f()
    for c in range(1, 4):
        ws3.cell(i, c).border = BORDER
    ws3.cell(i, 2).fill = INPUT_FILL
    ws3.cell(i, 3).fill = INPUT_FILL
    ws3.cell(i, 3).number_format = "yyyy/mm/dd"
    ws3.row_dimensions[i].height = 24
ws3.freeze_panes = "A5"
page_setup(ws3, "4:4", landscape=True)

# ---- 使い方・根拠法令 ------------------------------------------------------------
ws4 = wb.create_sheet("使い方・根拠法令")
ws4.column_dimensions["A"].width = 26
ws4.column_dimensions["B"].width = 44
ws4.column_dimensions["C"].width = 34
rows = [
    ("使い方", None, None, "h"),
    ("黄色のセル", "記入欄（判定・備考・氏名・日付など）", None, None),
    ("判定 ○", "適正", None, None),
    ("判定 ×", "不適。是正記録シートに転記する", None, None),
    ("判定 －", "該当なし（当日その作業がない）", None, None),
    (None, None, None, None),
    ("点検の頻度（目安）", None, None, "h"),
    ("区分", "頻度", "主な実施者", "th"),
    ("作業開始前点検", "毎日", "職長・作業主任者", None),
    ("定期点検", "週1回", "元請の現場代理人・安全担当", None),
    ("安全パトロール", "月1回", "店社の安全担当・協力会社", None),
    ("臨時点検", "悪天候・地震の後、足場や支保工の組立・変更の後", "作業主任者・点検者に指名された者", None),
    (None, None, None, None),
    ("是正の流れ", None, None, "h"),
    ("1", "×が付いた項目は、その場で作業を止めるかどうかを判断する", None, None),
    ("2", "是正記録に、内容・担当・期限を書く", None, None),
    ("3", "是正後に再点検し、確認者が署名する", None, None),
    (None, None, None, None),
    ("参考：主な根拠法令", None, None, "h"),
    ("法令", "主に関係する章", None, "th"),
] + [(a, b, None, None) for a, b in LAWS] + [
    (None, None, None, None),
    ("注意", "数値基準は主な規定を記憶に基づいて記載した案。採用前に最新の法令・発注者の仕様書・自社の安全基準と照合すること。", None, None),
]
for i, (a, b, c, kind) in enumerate(rows, 1):
    if kind == "h":
        ws4.cell(i, 1, a).font = f(True, 12)
        continue
    if kind == "th":
        labels = [x for x in (a, b, c) if x]
        for ci, lab in enumerate(labels, 1):
            cell = ws4.cell(i, ci, lab)
            cell.font = f(True, 10, "FFFFFF")
            cell.fill = HEAD_FILL
            cell.border = BORDER
            cell.alignment = CENTER
        continue
    for ci, v in enumerate((a, b, c), 1):
        if v is not None:
            cell = ws4.cell(i, ci, v)
            cell.font = f(color="C00000" if a == "注意" else "000000")
            cell.alignment = WRAP
    if a == "黄色のセル":
        ws4.cell(i, 1).fill = INPUT_FILL
    if a == "注意":
        ws4.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        ws4.row_dimensions[i].height = 32
for i in range(9, 13):
    for c in range(1, 4):
        ws4.cell(i, c).border = BORDER
law_start = [i for i, r_ in enumerate(rows, 1) if r_[0] == "法令"][0] + 1
for i in range(law_start, law_start + len(LAWS)):
    for c in range(1, 3):
        ws4.cell(i, c).border = BORDER
    ws4.row_dimensions[i].height = 30
page_setup(ws4)

wb.save(OUT)
print("wrote", OUT, "items:", sum(len(s[1]) for s in SECTIONS), "rows", first, last)
