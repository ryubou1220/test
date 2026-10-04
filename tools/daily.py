import sys, re, datetime, openpyxl
src=open('make_joubun.py').read(); src=src[:src.index('SRC=sys.argv')]
g={}; exec(src,g)
b=open('make_joubun.py').read()
# pull CL / classify / project / content / ROUTINE definitions
i=b.index("# 現場技術業務委託共通仕様書"); j=b.index("def build(")
exec(b[i:j],g)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L
ents,days,GNO,GNM=g['parse']('d9.xlsx')
ws0=openpyxl.load_workbook('d9.xlsx').active
starts=[c.row for c in ws0['H'] if c.value=='業務処理結果報告書']
WX={no:str(ws0.cell(s+19,23).value or '').strip() for s,(d,no) in zip(starts,days)}
T0,T1=datetime.date(2026,9,28),datetime.date(2026,9,30)
sel=[e for e in ents if T0<=datetime.date.fromisoformat(e['date'])<=T1]
Z=str.maketrans('０１２３４５６７８９：','0123456789:')
NM={'業務内容打合せ':'業務内容打合せ','業務打合せ':'翌日作業打合せ','『業務処理結果報告書』作成':'業務処理結果報告書作成'}
LOC={('2026-09-29','工事初回打合せ'):'実施場所：河川砂防課（確認済み）','2026-09-28':None}
FN='ＭＳ Ｐ明朝'
def F(**k): return Font(name=FN,size=k.pop('size',10),**k)
th=Side(style='thin'); BOX=Border(left=th,right=th,top=th,bottom=th)
HDR=PatternFill('solid',fgColor='D9E1F2'); SEC=PatternFill('solid',fgColor='EDEDED')
WR=Alignment(wrap_text=True,vertical='center'); CE=Alignment(horizontal='center',vertical='center',wrap_text=True)
wb=Workbook(); ws=wb.active; ws.title='日別業務一覧'
wm=wb.create_sheet('条文マスタ')
CL=g['CL']; N=len(CL)
for i,h in enumerate(['条文番号','章','見出し','条文要旨'],1):
    c=wm.cell(1,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CE
for r,row in enumerate(CL,2):
    for i,v in enumerate(row,1):
        c=wm.cell(r,i,v); c.font=F(); c.border=BOX; c.alignment=CE if i<=2 else WR
for i,w in enumerate([15,24,36,70],1): wm.column_dimensions[L(i)].width=w
MK=f"条文マスタ!$A$2:$A${1+N}"; MC=f"条文マスタ!$C$2:$C${1+N}"

ws['A1']='日別業務一覧（条文対応版）　令和8年9月分（9月28日～9月30日）'; ws['A1'].font=F(size=14,bold=True)
ws['A2']=f'業務番号：{GNO}　　業務名：{GNM}'; ws['A2'].font=F(size=9)
ws['A3']='出典：業務処理結果報告書 №120～№122。条文は現場技術業務委託共通仕様書（京都府）。'; ws['A3'].font=F(size=9)
H=['No.','開始時刻','業務区分','実施内容','対象工事','相手先','担当監督員','条文番号','見出し','備考']
HR=5
for i,h in enumerate(H,1):
    c=ws.cell(HR,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CE
r=HR; cur=None; n=0
for e in sel:
    if e['date']!=cur:
        cur=e['date']; d=datetime.date.fromisoformat(cur); r+=1
        ws.cell(r,1,f"{d.month}月{d.day}日（{'月火水木金土日'[d.weekday()]}）　業務処理結果報告書 №{e['no']}　天候：{WX[e['no']]}　勤務時間 8:30～17:30")
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=10)
        for i in range(1,11): c=ws.cell(r,i); c.fill=SEC; c.border=BOX; c.font=F(bold=True)
        n=0
    r+=1; n+=1
    routine=e['cat'] in g['ROUTINE']
    key=g['classify'](e)
    cat='執務室業務' if e['cat'].startswith('執務室') else ('打合せ・報告' if routine else ('河川砂防課' if e['cat']=='河川砂防課' else '工事関係'))
    txt=NM[e['cat']] if routine else g['content'](e)
    tanto=e['tanto'].replace(' 翌日作業打合せ','')
    note=''
    if '初回打合せ' in e['det'] and cur=='2026-09-29': note='実施場所：河川砂防課（確認済み）'
    if '完成図書下検査' in e['det']: note='実施場所：河川砂防課'
    vals=[n,e['time'].replace('～','').translate(Z),cat,txt,'―' if routine else g['project'](e),
          ' / '.join(x for x in [e['insp'],e['aite']] if x) or '―',tanto or '―',key,f'=INDEX({MC},MATCH(H{r},{MK},0))',note]
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.font=F(color='0000FF' if i==8 else '000000'); c.border=BOX
        c.alignment=CE if i in (1,2,3,7,8) else WR
    ws.row_dimensions[r].height=30
last=r
# summary
r+=2; ws.cell(r,1,'条文別件数（3日間）').font=F(bold=True,size=11); r+=1
keys=[k[0] for k in CL if any(g['classify'](e)==k[0] for e in sel)]
day_rows={}; cd=None
for row in range(HR+1,last+1):
    v=ws.cell(row,1).value
    if isinstance(v,str) and '月' in v: cd=v[:v.index('日')+1]; day_rows[cd]=[row+1,row+1]
    else: day_rows[cd][1]=row
# layout: A:C 条文番号 / D 見出し / E,F,G 日別 / H 計
heads=[(1,'条文番号'),(4,'見出し')]+[(5+j,dk.replace('月','/').replace('日','')) for j,dk in enumerate(day_rows)]+[(8,'計')]
def band(rr,fill=None,bold=False):
    ws.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=3)
    for i in range(1,9):
        c=ws.cell(rr,i); c.border=BOX; c.font=F(bold=bold); c.alignment=CE if i!=4 else WR
        if fill: c.fill=fill
for col,h in heads: ws.cell(r,col,h)
band(r,HDR,True)
s0=r+1
for k in keys:
    r+=1; ws.cell(r,1,k); ws.cell(r,4,f'=INDEX({MC},MATCH(A{r},{MK},0))')
    for j,(dk,(a_,b_)) in enumerate(day_rows.items()):
        ws.cell(r,5+j,f'=COUNTIF($H${a_}:$H${b_},$A{r})')
    ws.cell(r,8,f'=SUM(E{r}:G{r})'); band(r)
r+=1; ws.cell(r,1,'合計')
for col in 'EFGH': ws.cell(r,'ABCDEFGH'.index(col)+1,f'=SUM({col}{s0}:{col}{r-1})')
band(r,SEC,True)
r+=2
for t in ['※ 日常定型業務（業務内容打合せ・翌日作業打合せ・報告書作成）も1件ずつ記載し、第5条（業務処理予定表及び結果報告書）に対応させた。',
          '※ H列「条文番号」（青字）を書き換えると、見出しと条文別件数に反映される。']:
    ws.cell(r,1,t).font=F(size=9); r+=1
for i,w in enumerate([5,8,11,40,30,14,13,14,30,24],1): ws.column_dimensions[L(i)].width=w
ws.freeze_panes=f'A{HR+1}'; ws.print_title_rows=f'{HR}:{HR}'
ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=9; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.oddFooter.center.text='&P / &N'
wm.page_setup.orientation='landscape'; wm.page_setup.paperSize=9; wm.page_setup.fitToWidth=1; wm.page_setup.fitToHeight=1; wm.sheet_properties.pageSetUpPr.fitToPage=True
out='/home/user/test/日別業務一覧_条文対応版_令和8年9月分_28-30日.xlsx'; wb.save(out); print(out)
for e in sel: print(e['date'][5:],e['time'],g['classify'](e),'|',e['cat'],'|',e['det'])
