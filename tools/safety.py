import json, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L
ents=json.load(open('field.json'))
FN='ＭＳ Ｐ明朝'
def F(**k): return Font(name=FN,size=k.pop('size',10),**k)
th=Side(style='thin'); BOX=Border(left=th,right=th,top=th,bottom=th)
HDR=PatternFill('solid',fgColor='D9E1F2'); SEC=PatternFill('solid',fgColor='EDEDED')
WR=Alignment(wrap_text=True,vertical='center'); CE=Alignment(horizontal='center',vertical='center',wrap_text=True)
Z=str.maketrans('０１２３４５６７８９：','0123456789:')

TYPES=[
 ('A','河口・海岸の定期観測・現況確認',
  '・波浪、高潮、離岸流による水際での転落・流され\n・砂浜、消波ブロック、岩場、濡れた護岸での転倒\n・強風・落雷\n・熱中症、単独作業時の急病',
  '・出発前に気象警報、波浪・潮位（潮汐表）、風速を確認し、高波・強風・雷のおそれがあれば中止・延期する\n・波打ち際、消波ブロック上、濡れた護岸には立ち入らない（観測は安全な位置から行う）\n・水際に近づく場合はライフジャケットを着用する\n・滑りにくい靴を履き、両手が空く装備（ザック等）とする\n・原則2人以上で行動し、行先・帰着予定を事務所に連絡する\n・退避経路と車両の駐車位置（波・浸水の届かない場所）を事前に確認する'),
 ('B','河川工事現場の立会・進捗確認・竣工検査（現場）',
  '・稼働中の重機（バックホウ旋回範囲）、ダンプ等の工事車両との接触\n・法面、護岸、仮設通路での転倒・転落\n・河川内作業区域での急な増水\n・熱中症',
  '・現場到着時に現場代理人へ連絡し、受注者の安全指示・誘導に従う（新規入場時は受注者の安全教育・KY活動に参加）\n・保護帽、安全靴、反射ベストを着用する\n・重機の旋回範囲、ダンプの走行経路に立ち入らない。近づく必要があるときはオペレーター・誘導員に合図して作業を止めてもらう\n・法肩、水際、仮締切の上には不用意に近づかない\n・上流の降雨・河川水位を確認し、増水のおそれがあれば河川内に立ち入らない\n・身分証明書を携帯する（特記仕様書「現場への立入り等」）'),
 ('C','残土仮置場・残土処理場の確認立会',
  '・ダンプの出入り、重機の敷均し作業との接触\n・盛土法面の崩壊、軟弱な盛土上での足元の不安定\n・粉じん',
  '・出入口でダンプの動線と誘導員の配置を確認し、指示に従う\n・盛土の上や法肩に乗らず、締固めの済んだ箇所から確認する\n・保護帽、安全靴、反射ベストを着用する\n・敷地内では車両を指定の場所に停め、徒歩移動は重機の死角を避ける'),
 ('D','既設井堰の現況調査（宇川）',
  '・堰体上や堰直下での転倒・転落、苔による滑り\n・流水中の歩行、急な増水\n・ハチ、ヘビ、ツキノワグマ等の野生生物\n・水田・水路の私有地への立入り',
  '・堰体の上を歩かない。流水中に入る場合は胴長・ライフジャケットを着用し、水深・流速を確認する\n・2人以上で行動し、1人は岸から見守る\n・事前に地元・水利関係者へ調査日時を連絡し、私有地は了解を得てから立ち入る\n・クマ鈴を携帯し、ハチの巣に注意する。長袖・長ズボン・手袋を着用する\n・前日・当日の降雨と河川水位を確認する'),
 ('E','砂防工事の現場調査（渓流部）',
  '・急斜面、転石、落石\n・渓流での転倒・増水\n・ツキノワグマ、ハチ等の野生生物',
  '・受注者と同行し、調査範囲・経路を事前に打ち合わせる\n・保護帽を着用し、斜面の下側に人がいるときは上で行動しない\n・クマ鈴を携帯し、単独行動をしない\n・降雨後・降雨中は渓流内に立ち入らない'),
]
COMMON=('共通','全現場共通',
 '・車両での移動中の交通事故\n・路肩駐車時の後続車との接触\n・熱中症（7～9月）\n・雨天時のスリップ、落雷',
 '・行先・同行者・帰着予定を事務所の行動予定に記入し、携帯電話で連絡が取れるようにする\n・路肩に駐車する場合はハザードランプ・三角表示板を使用し、車道側に出ない\n・WBGT（暑さ指数）を確認し、水分・塩分を補給し、こまめに休憩する。体調不良時は作業を中止する\n・雷鳴が聞こえたら直ちに車内等へ退避する\n・保護帽、安全靴を常備し、身分証明書を携帯する')
TYP={t[0]:t for t in TYPES}
SHORT={
 'A':('水際転落・高波、転倒、熱中症','気象・波浪・潮位を事前確認／波打ち際・消波ブロックに立ち入らない／水際ではライフジャケット着用／2人以上で行動'),
 'B':('重機・工事車両との接触、法面転落、増水','現場代理人へ連絡し受注者の安全指示に従う／保護帽・安全靴・反射ベスト着用／重機旋回範囲に入らない／水位確認'),
 'C':('ダンプ・重機との接触、盛土法面崩壊','ダンプ動線と誘導員の指示に従う／盛土上・法肩に乗らない／保護帽・安全靴・反射ベスト着用'),
 'D':('堰体上での転倒・転落、増水、野生生物','堰体上を歩かない／流水中は胴長・ライフジャケット着用／2人以上で行動／地元・水利関係者へ事前連絡／クマ鈴携帯'),
 'E':('急斜面・落石、渓流での転倒、野生生物','受注者と同行し経路を事前確認／保護帽着用・上下作業を避ける／クマ鈴携帯・単独行動禁止'),
}


def typ(e):
    s=e['cat']+e['det']
    if re.search('定期観測|河口現況確認',s): return 'A'
    if '井堰' in s: return 'D'
    if '残土' in s: return 'C'
    if '現場調査' in s and '成願寺' in e['cat']: return 'E'
    if re.search('確認立会|現場立会|進捗状況確認|現場状況確認|竣工検査立会（現',s): return 'B'
    return None
def proj(e):
    c=e['cat'].replace('　',' ')
    if re.search('定期観測|河口現況確認',c): return c.replace(' 定期観測箇所','').replace(' 定期観測','').replace('現況確認','')+' 定期観測'
    if '井堰' in c: return '宇川 現況既設井堰調査'
    return c
def place(e):
    d=e['det']; m=re.search(r'[（(]([^）)]*)[）)]',d)
    s=e['cat']+d
    if '網野町浜詰' in d: return '網野町浜詰'
    if '木津川河口' in s: return '網野町（木津川河口）'
    if '宇川河口' in s and '久僧' in s: return '丹後町上野・久僧'
    if '宇川河口' in s: return '丹後町上野（宇川河口）'
    if '久僧' in s: return '丹後町久僧（久僧海岸・吉野川河口）'
    if '新樋越川' in s: return '新樋越川河口'
    if '井堰' in s: return d.replace('丹後・宇川 ','').replace('弥栄・宇川 ','')
    if '金谷' in d: return '久美浜町金谷・畑地区'
    if m and '・' in m[1]: return m[1].replace('  ',' ')
    return '―'
def content(e):
    d=e['det']
    if typ(e)=='A': return '観測調査作業' if '観測' in d else '定期観測箇所の現況確認'
    if '井堰' in e['cat']: return '既設井堰の現況調査（'+re.search(r'(\S+箇所)',d)[1]+'）' if '箇所' in d else '既設井堰の現況調査'
    if '金谷' in d: return '現場立会'
    return re.sub(r'\s*[（(][^）)]*[）)]','',d).replace(' 倉垣副主査','').strip() or d
def note(e,t):
    d=datetime.date.fromisoformat(e['date']); n=[]
    if '雨' in e['wx']: n.append('雨天：足元の滑り、河川・水路の増水に注意。観測・立会の実施可否を判断')
    if d.month in (7,8) or (d.month==9 and d.day<=20 and '晴' in e['wx']): n.append('夏季：熱中症対策（WBGT確認・水分塩分補給・休憩）')
    if t=='B' and '竣工検査' in e['det']: n.append('検査官・監督員と同行。完成箇所の法面・護岸上の移動に注意')
    if e['time'].startswith('PM11'): n.append('※日報の開始時刻「PM11:30」はAM11:30の誤記と思われる')
    return '\n'.join('・'+x for x in n)

rows=[]
ref=[]
for e in ents:
    t=typ(e)
    s=e['cat']+e['det']
    if t: rows.append((e,t))
    elif re.search('下検査（主任|着手前打合せ|初回打合せ',s) and '入札室' not in s:
        ref.append(e)
def tk(t):
    t=t.replace('～','').translate(Z).replace('PM11','AM11'); m=re.match(r'(AM|PM)(\d+):(\d+)',t)
    return (int(m[2])%12+(12 if m[1]=='PM' else 0))*60+int(m[3])
rows.sort(key=lambda x:(x[0]['date'],tk(x[0]['time'])))

wb=Workbook()
ws=wb.active; ws.title='現場確認日一覧'
wt=wb.create_sheet('安全対策（類型別）'); wa=wb.create_sheet('月別集計'); wr=wb.create_sheet('参考（現地か不明）')

# --- 類型別
wt['A1']='現場確認業務の類型別 想定される危険と安全対策'; wt['A1'].font=F(size=14,bold=True)
wt['A2']='業務処理結果報告書（令和8年7～9月）の現場業務から類型を整理し、各類型で考えられる安全対策をまとめたもの（日報に記載された実施記録ではない）。'; wt['A2'].font=F(size=9)
for i,h in enumerate(['類型','業務の種類','想定される主な危険','考えられる安全対策'],1):
    c=wt.cell(4,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CE
for r,t in enumerate(TYPES+[COMMON],5):
    for i,v in enumerate(t,1):
        c=wt.cell(r,i,v); c.font=F(bold=(i<=2)); c.border=BOX; c.alignment=CE if i==1 else WR
        if t[0]=='共通': c.fill=SEC
    wt.row_dimensions[r].height=15*max(t[2].count('\n')+1,t[3].count('\n')+1+len(t[3])//60)
for i,w in enumerate([7,24,40,80],1): wt.column_dimensions[L(i)].width=w
TR=f"'安全対策（類型別）'!$A$5:$A${4+len(TYPES)}"
def tl(col,cell): return f"=INDEX('安全対策（類型別）'!${col}$5:${col}${4+len(TYPES)},MATCH({cell},{TR},0))"

# --- 一覧
ws['A1']='現場確認日一覧及び考えられる安全対策（令和8年7月～9月）'; ws['A1'].font=F(size=14,bold=True)
ws['A2']='出典：業務処理結果報告書（令和8年7月分 №62～83、8月分 №84～103、9月分 №104～122）。執務室業務・河川砂防課での打合せ・書面検査は除外。安全対策は類型ごとに考えられるもの（日報に実施記録はない）。'; ws['A2'].font=F(size=9)
H=['No.','実施日','曜','天候','報告書№','開始時刻','類型','工事・業務名','場所','内容','担当監督員','受注者等','想定される主な危険','考えられる主な安全対策（詳細は類型別シート）','当日の留意点']
HR=4
for i,h in enumerate(H,1):
    c=ws.cell(HR,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CE
r=HR; prevm=None; n=0
for e,t in rows:
    d=datetime.date.fromisoformat(e['date'])
    if d.month!=prevm:
        r+=1; ws.cell(r,1,f'令和8年{d.month}月')
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=15)
        for i in range(1,16): c=ws.cell(r,i); c.fill=SEC; c.border=BOX; c.font=F(bold=True)
        prevm=d.month
    r+=1; n+=1
    vals=[n,d,f'=MID("日月火水木金土",WEEKDAY(B{r}),1)',e['wx'],e['no'],e['time'].replace('～','').translate(Z),t,
          f"=G{r}&\"：\"&{tl('B',f'G{r}')[1:]}",None,None,e['tanto'].replace('  ',' '),e['aite'] or '―',
          SHORT[t][0],SHORT[t][1],note(e,t)]
    vals[7]=proj(e); vals[8]=place(e); vals[9]=content(e)
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.font=F(size=9); c.border=BOX
        c.alignment=CE if i in (1,2,3,4,5,6,7,11) else Alignment(wrap_text=True,vertical='top')
    ws.cell(r,2).number_format='m"月"d"日"'; ws.cell(r,5).number_format='"№"0'
    ws.row_dimensions[r].height=48
last=r
for i,w in enumerate([5,8,4,9,7,8,5,28,18,18,11,13,24,50,34],1): ws.column_dimensions[L(i)].width=w
ws.freeze_panes=f'A{HR+1}'; ws.print_title_rows=f'{HR}:{HR}'
ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=9
ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
ws.oddFooter.center.text='&P / &N'
ws.auto_filter.ref=f'A{HR}:O{last}'

# --- 月別集計
wa['A1']='月別 現場確認の実施状況（令和8年7月～9月）'; wa['A1'].font=F(size=14,bold=True)
hd=['月','履行日数','現場確認日数','現場確認の割合','現場確認件数']+[f'類型{t[0]}' for t in TYPES]
for i,h in enumerate(hd,1):
    c=wa.cell(3,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CE
B=f"'現場確認日一覧'!$B${HR+1}:$B${last}"; G=f"'現場確認日一覧'!$G${HR+1}:$G${last}"
for k,(m,days) in enumerate([(7,22),(8,20),(9,19)]):
    rr=4+k
    s0=f'DATE(2026,{m},1)'; s1=f'DATE(2026,{m+1},1)'
    vals=[f'{m}月',days,
          f'=SUMPRODUCT(({B}>={s0})*({B}<{s1})/COUNTIF({B},{B}&""))',
          f'=IF(B{rr}=0,0,C{rr}/B{rr})',
          f'=COUNTIFS({B},">="&{s0},{B},"<"&{s1})']+[f'=COUNTIFS({B},">="&{s0},{B},"<"&{s1},{G},"{t[0]}")' for t in TYPES]
    for i,v in enumerate(vals,1):
        c=wa.cell(rr,i,v); c.font=F(color='0000FF' if i==2 else '000000'); c.border=BOX; c.alignment=CE
    wa.cell(rr,4).number_format='0.0%'
rr=7; wa.cell(rr,1,'合計')
for i in range(2,len(hd)+1):
    col=L(i); wa.cell(rr,i,f'=SUM({col}4:{col}6)' if i!=4 else '=IF(B7=0,0,C7/B7)')
    wa.cell(rr,i).number_format='0.0%' if i==4 else 'General'
for i in range(1,len(hd)+1):
    c=wa.cell(rr,i); c.font=F(bold=True); c.fill=SEC; c.border=BOX; c.alignment=CE
notes=['※ 履行日数（青字）は業務処理結果報告書の枚数（7月22日、8月20日、9月19日）。',
       '※ 現場確認日数は「現場確認日一覧」の実施日の重複を除いた日数。件数は一覧の行数。']
for i,t in enumerate(TYPES): notes.append(f'※ 類型{t[0]}：{t[1]}')
for i,t in enumerate(notes): wa.cell(9+i,1,t).font=F(size=9)
for i,w in enumerate([8,9,12,13,12,8,8,8,8,8],1): wa.column_dimensions[L(i)].width=w

# --- 参考
wr['A1']='参考：現地で実施した可能性があるが、日報からは判断できない業務（一覧・集計には含めていない）'; wr['A1'].font=F(size=12,bold=True)
for i,h in enumerate(['実施日','報告書№','開始時刻','工事・業務名','内容','担当監督員','受注者等','除外理由'],1):
    c=wr.cell(3,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CE
why={'下検査':'工事完成下検査は書面中心と考えられるが、現地確認を含む可能性がある','着手前':'打合せ場所の記載がない','初回':'打合せ場所の記載がない'}
for k,e in enumerate(sorted(ref,key=lambda e:e['date']),4):
    s=e['cat']+e['det']; reason=next(v for kk,v in why.items() if kk in s)
    vals=[datetime.date.fromisoformat(e['date']),e['no'],e['time'].replace('～','').translate(Z),e['cat'],e['det'],e['tanto'],e['aite'] or '―',reason]
    for i,v in enumerate(vals,1):
        c=wr.cell(k,i,v); c.font=F(size=9); c.border=BOX; c.alignment=CE if i<=3 else WR
    wr.cell(k,1).number_format='m"月"d"日"'; wr.cell(k,2).number_format='"№"0'; wr.row_dimensions[k].height=30
for i,w in enumerate([9,8,8,34,26,11,13,40],1): wr.column_dimensions[L(i)].width=w
for w_ in (wt,wa,wr):
    w_.page_setup.orientation='landscape'; w_.page_setup.paperSize=9; w_.page_setup.fitToWidth=1; w_.page_setup.fitToHeight=0; w_.sheet_properties.pageSetUpPr.fitToPage=True
wt.page_setup.fitToHeight=1
out='/home/user/test/現場確認日_安全対策一覧_7-9月分.xlsx'; wb.save(out)
print(out,n,'rows; ref',len(ref))
import collections
print(collections.Counter((datetime.date.fromisoformat(e['date']).month) for e,t in rows))
print({m:len({e['date'] for e,t in rows if e['date'][5:7]==f'{m:02d}'}) for m in (7,8,9)})
