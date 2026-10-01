#!/usr/bin/env python3
"""現場技術業務 検査受検資料 条文別整理表 作成ツール

使い方:  python make_joubun.py <業務処理結果報告書(月分).xlsx> [出力フォルダ]
  - 日報（業務処理結果報告書）1か月分のExcelを読み込み、目次・条文別整理表・
    条文別集計・条文マスタの4シートを持つブックを作成する。
  - 条文は「現場技術業務委託共通仕様書（京都府）」に準拠。振分けルールは classify() を参照。
  - LibreOffice(soffice) がある場合は、整理表の印刷頁を求めて目次に記入する。
"""
import sys, os, json, datetime, re, subprocess, tempfile, shutil
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter as L

def parse(path):
    ws=openpyxl.load_workbook(path).active
    starts=[c.row for c in ws['H'] if c.value=='業務処理結果報告書']
    if not starts: sys.exit('業務処理結果報告書の様式が見つかりません: '+path)
    ents=[]; d=None; days=[]
    first_no=ws.cell(starts[0]+2,37).value
    if not isinstance(first_no,int): sys.exit('先頭の報告書№（AK列）が数値ではありません')
    gyomu_no=str(ws.cell(starts[0]+4,19).value or '').strip(); gyomu_nm=str(ws.cell(starts[0]+6,19).value or '').strip()
    for i,s in enumerate(starts):
        dr=s+17; dv=ws.cell(dr,4).value
        if isinstance(dv,datetime.datetime): d=dv.date()
        elif isinstance(dv,str) and re.match(r'=D\d+\+\d+$',dv) and d: d=d+datetime.timedelta(days=int(dv.split('+')[1]))
        else: sys.exit(f'{dr}行目の日付を読み取れません: {dv!r}')
        rno=first_no+i; days.append((d,rno))
        end=starts[i+1] if i+1<len(starts) else ws.max_row+1
        g=lambda rr,cc: (ws.cell(rr,cc).value or '')
        for r in range(s,end):
            t=ws.cell(r,3).value
            if isinstance(t,str) and '～' in t:
                ents.append(dict(no=rno,date=d.isoformat(),time=t.strip(),
                  cat=str(g(r+1,4)).strip(),tanto=str(g(r+1,28)).replace('担当','').replace('：','').replace('　',' ').strip(),
                  det=' '.join(str(x).strip() for x in [g(r+2,5),g(r+2,12)] if x).replace('　',' ').strip(),
                  insp=str(g(r+2,13)).strip(), aite=str(g(r+2,22) or g(r+1,22)).strip()))
    return ents, days, gyomu_no, gyomu_nm

SRC=sys.argv[1] if len(sys.argv)>1 else sys.exit(__doc__)
OUTDIR=sys.argv[2] if len(sys.argv)>2 else '.'
ents,DAYS,GNO,GNM=parse(SRC)
WD='月火水木金土日'
d0,dl=DAYS[0][0],DAYS[-1][0]
Y,M=d0.year,d0.month; RY=Y-2018
mstart=datetime.date(Y,M,1); mend=(datetime.date(Y+M//12,M%12+1,1)-datetime.timedelta(days=1))
NDAYS=len(DAYS); NO0,NO1=DAYS[0][1],DAYS[-1][1]
TAG=f'令和{RY}年{M}月分'
jd=lambda d: f'令和{d.year-2018}年{d.month}月{d.day}日（{WD[d.weekday()]}）'
PERIOD=f'{jd(mstart)}～ {jd(mend)}　履行日数 {NDAYS}日（業務処理結果報告書 №{NO0}～№{NO1}）'
FN='ＭＳ Ｐ明朝'
def F(**k): return Font(name=FN, size=k.pop('size',10), **k)
thin=Side(style='thin'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
HDR=PatternFill('solid',fgColor='D9E1F2'); SEC=PatternFill('solid',fgColor='EDEDED'); NONE=PatternFill('solid',fgColor='F7F7F7')
WRAP=Alignment(wrap_text=True,vertical='center'); CEN=Alignment(horizontal='center',vertical='center',wrap_text=True)

# 現場技術業務委託共通仕様書（京都府）
CL=[
 ('第3条第3項(1)','第1章 総則','管理技術者等（監督に関する業務の厳正な実施）','必要な監督に関する業務を厳正に実施すること。'),
 ('第3条第3項(4)','第1章 総則','管理技術者等（工事請負者等への連絡・通知）','工事請負者又は外部への連絡若しくは通知を行う場合には、その内容を相手に正確に伝えること。'),
 ('第3条第3項(5)','第1章 総則','管理技術者等（設計図書の理解・現場状況の精通）','請負工事の契約書及び設計図書等の内容を十分理解し、更に工事現場の状況についても精通しておくこと。'),
 ('第3条第3項(6)','第1章 総則','管理技術者等（図書の整理）','業務の実施にあたっては、業務に関する図書を適切に整理しておくこと。'),
 ('第3条第4項','第1章 総則','管理技術者等（監督職員との打合せ・協議）','管理技術者は、特記仕様書に定めるところにより監督職員と打合せ・協議を行い、その結果について相互に確認する（特記仕様書：着手時及び最終時の計2回）。'),
 ('第5条','第1章 総則','業務処理予定表及び結果報告書','業務処理予定表及び業務処理結果報告書（実施した業務の内容、その他必要事項）を作成し、監督職員に提出する。'),
 ('第6条','第1章 総則','業務完了時の提出書類','業務完了時に、業務処理結果報告書を一括整理して提出する。'),
 ('第7条','第2章 監督に関する現場技術業務','書類の確認','工事請負者から提出された書類は、これを確認し監督職員に報告しなければならない。'),
 ('第8条','第2章 監督に関する現場技術業務','立会','設計図書に基づき立会いを行った場合は、その結果を書面で監督職員に報告する。'),
 ('第9条','第2章 監督に関する現場技術業務','検測','請負工事の施工が設計図書に示す所定の出来形及び品質を確保するため現地で検測を行い、その成果を監督職員に提出する。'),
 ('第10条','第2章 監督に関する現場技術業務','材料検査','材料検査を実施したときは、検査年月日・品名・寸法等・検査数量・検査結果及び合格数量等を付記して監督職員に提出する。'),
 ('第11条','第2章 監督に関する現場技術業務','工程管理','請負工事の進捗状況を把握し、工事が遅延するおそれがあれば遅滞なく書面で監督職員に報告する。'),
 ('第12条','第2章 監督に関する現場技術業務','品質管理','工事受注者が仕様書に定められた品質管理試験を忠実に実行しているか確認し、その結果を書面で報告する。'),
 ('第13条','第2章 監督に関する現場技術業務','図面と現地の不一致等','設計図書と工事現場の状態の不一致等の通知を受けたときは、遅滞なく書面で監督職員に報告する。'),
 ('第14条','第2章 監督に関する現場技術業務','検査の立会','請負工事に係る工事検査及び監督職員が行う検査に立会い、求められる説明に応じる。'),
 ('第15条','第2章 監督に関する現場技術業務','設計変更工事検査等に関する図書','監督職員と協議のうえ、設計変更、工事完成検査若しくは既済部分検査等に必要な測量、測定又は図書等の資料作成を行う。'),
 ('第16条','第2章 監督に関する現場技術業務','対外接渉に関する資料','監督職員と協議の上、地元若しくは関係機関等との接渉に必要な測量、調査又は資料の作成を行う。'),
 ('第17条','第2章 監督に関する現場技術業務','書面での報告','第2章の各条にいう書面での監督職員に対する報告は、業務処理結果報告書による。'),
 ('第19条','第3章 設計に関する現場技術業務','設計に必要な調査','設計に必要な現場条件等の調査を、事前に監督職員と協議の上行い、調査結果は書面で提出する。'),
 ('第20条','第3章 設計に関する現場技術業務','設計に必要な資料','設計に必要な図面、数量取りまとめ、各種データの作成等を、事前に監督職員と協議の上行い、結果を書面で提出する。'),
]
ORDER={c[0]:i for i,c in enumerate(CL)}
UNK=[]
ROUTINE=('業務内容打合せ','業務打合せ','『業務処理結果報告書』作成')
def classify(e):
    s=e['cat']+' '+e['det']
    if e['cat'] in ROUTINE or '社内資料' in s: return '第5条'
    if '検査資料取り纏め' in s: return '第3条第3項(6)'
    if '定期観測' in s: return '第19条'
    if '過積載' in s: return '第3条第3項(1)'
    if '竣工検査' in s or e['cat'].startswith('工事完成下検査') or '完成図書下検査' in s: return '第14条'
    if re.search('成果品|プロセスチェック|出来形図書|下検査指摘|数量|土量',s): return '第15条'
    if '初回打合せ' in s: return '第3条第3項(4)'
    if re.search('地元立会|案内ビラ|着手前|参考計画書',s): return '第16条'
    if re.search('現場立会|立会議事|立会メモ|追記資料',s): return '第8条'
    if re.search('設計図書|打合せ事項メモ|現場調査',s): return '第3条第3項(5)'
    if re.search('施工計画書|施工体制台帳|再生資源|ASP|修正資料',s): return '第7条'
    UNK.append(f"№{e['no']} {e['date']} {e['time']} {s}"); return '第5条'

def project(e):
    s=e['cat']+' '+e['det']+' '+e['aite']
    if '定期観測' in s: return '河口・海岸 定期観測'
    if '小西川' in s: return '竹野川（小西川）広域河川改修（補正・防災安全）工事'+('ほか7件' if 'ほか7件' in s else '')
    if '宇川' in s: return '管内一円（宇川）府民協働型インフラ保全工事他３件'
    if '川上谷川' in s and '成願寺川' in s: return '川上谷川緊急浚渫推進（河川）工事／管内一円（成願寺川）緊急浚渫推進（砂防）工事'
    if '川上谷川' in s or '小野澤' in s: return '川上谷川緊急浚渫推進（河川）工事'
    if '成願寺川' in s or 'サンキ' in s: return '管内一円（成願寺川）緊急浚渫推進（砂防）工事'+('他' if '成願寺川他' in s else '')
    if '河梨川' in s or '西田' in s: return '管内一円（河梨川他）府民協働型インフラ保全工事'
    if '竹野川他' in s or '修己' in s: return '竹野川他府民協働型インフラ保全工事'
    if '過積載' in s: return '第２四半期調査'
    if '検査資料' in s: return '現場技術業務'
    return '―'

def content(e):
    if e['det']:
        if e['cat'] in ('執務室業務','書面資料打合せ','河川砂防課') or e['cat'].startswith('執務室業務'):
            return {'書面資料打合せ':'【書面資料打合せ】','河川砂防課':'【河川砂防課】'}.get(e['cat'],'')+e['det']
        if '定期観測' in e['cat']: return e['cat'].replace('　',' ')+'：'+e['det']
        if e['cat'].startswith('工事完成下検査'): return e['cat']
        return e['det']
    return e['cat']
def build(PAGES_IN):
    Z=str.maketrans('０１２３４５６７８９：','0123456789:')
    rows=[]; by_day={}
    for e in ents:
        c=classify(e)
        if e['cat'] in ROUTINE: by_day.setdefault(e['date'],[]).append(e); continue
        rows.append(dict(c=c,date=e['date'],time=e['time'].replace('～','').translate(Z),proj=project(e),
            txt=content(e),aite=' / '.join(x for x in [e['insp'],e['aite']] if x),tanto=e['tanto'],no=e['no'],note=''))
    NM={'業務内容打合せ':'業務内容打合せ（業務処理予定の確認）','業務打合せ':'翌日作業打合せ（業務処理予定）','『業務処理結果報告書』作成':'業務処理結果報告書作成'}
    for d,t in by_day.items():
        rows.append(dict(c='第5条',date=d,time=t[0]['time'].replace('～','').translate(Z),proj='―',
            txt='、'.join(f"{x['time'].replace('～','').translate(Z)} {NM[x['cat']]}" for x in t),aite='',
            tanto=next((x['tanto'] for x in t if x['cat']=='業務内容打合せ'),''),no=t[0]['no'],note='日常定型業務（1日1行に集約）'))
    def tkey(t):
        m=re.match(r'(AM|PM)(\d+):(\d+)',t); return (int(m[2])%12+(12 if m[1]=='PM' else 0))*60+int(m[3])
    rows.sort(key=lambda r:(ORDER[r['c']],r['date'],tkey(r['time']),r['txt']))
    if UNK: print('【要確認】振分けルールに該当しない業務（第5条に仮置き）:\n  '+'\n  '.join(UNK))

    wb=Workbook(); wm=wb.active; wm.title='条文マスタ'
    ws=wb.create_sheet('条文別整理表',0); wa=wb.create_sheet('条文別集計',1)
    N=len(CL); MK=f"条文マスタ!$A$5:$A${4+N}"
    def mref(col): return f"条文マスタ!${col}$5:${col}${4+N}"
    def look(col,cell): return f'=INDEX({mref(col)},MATCH({cell},{MK},0))'

    # 条文マスタ
    wm['A1']='条文マスタ（現場技術業務委託共通仕様書 条文一覧）'; wm['A1'].font=F(size=14,bold=True)
    wm['A2']='出典：現場技術業務委託共通仕様書（京都府、令和元年8月20日 元指第486号 最終改正）別添2 及び 現場技術業務 特記仕様書。条文要旨は原文を要約したもの。'; wm['A2'].font=F(size=9)
    for i,h in enumerate(['条文番号','章','見出し','条文要旨'],1):
        c=wm.cell(4,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CEN
    for r,row in enumerate(CL,5):
        for i,v in enumerate(row,1):
            c=wm.cell(r,i,v); c.font=F(); c.border=BOX; c.alignment=CEN if i<=2 else WRAP
        wm.row_dimensions[r].height=32
    for i,w in enumerate([15,24,36,70],1): wm.column_dimensions[L(i)].width=w
    wm.page_setup.orientation='landscape'; wm.page_setup.paperSize=9; wm.page_setup.fitToWidth=1; wm.page_setup.fitToHeight=0; wm.sheet_properties.pageSetUpPr.fitToPage=True

    # 条文別整理表
    ws['A1']='現場技術業務　検査受検資料　条文別整理表（'+TAG+'）'; ws['A1'].font=F(size=16,bold=True)
    ws.merge_cells('A1:L1'); ws['A1'].alignment=Alignment(horizontal='center')
    info=[('業務番号',GNO),('業務名',GNM),
          ('対象期間',PERIOD),
          ('適用仕様書','現場技術業務委託共通仕様書（京都府）及び 現場技術業務 特記仕様書')]
    for i,(k,v) in enumerate(info,3):
        ws.cell(i,1,k).font=F(bold=True); ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=2)
        ws.cell(i,3,v).font=F(); ws.merge_cells(start_row=i,start_column=3,end_row=i,end_column=12)
    H=['条文番号','見出し','No.','実施日','曜','開始時刻','対象工事・業務','実施内容','相手先／検査官','担当監督員','報告書№','備考']
    HR=8
    for i,h in enumerate(H,1):
        c=ws.cell(HR,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CEN
    r=HR+1
    HROWS={}
    for key,ch,title,summ in CL:
        sub=[x for x in rows if x['c']==key]; n=max(len(sub),1); a,b=r+1,r+n
        HROWS[key]=r
        ws.cell(r,1,key); ws.cell(r,2,look('C',f'$A{r}'))
        ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6)
        ws.cell(r,7,look('D',f'$A{r}')); ws.merge_cells(start_row=r,start_column=7,end_row=r,end_column=8)
        ws.cell(r,9,f'=COUNT($C${a}:$C${b})&"件"'); ws.cell(r,10,'実施日数')
        ws.cell(r,11,f'=IF(COUNT($C${a}:$C${b})=0,0,SUMPRODUCT(($C${a}:$C${b}<>"")/COUNTIFS($D${a}:$D${b},$D${a}:$D${b}&"")))&"日"')
        for i in range(1,13):
            c=ws.cell(r,i); c.fill=SEC; c.font=F(bold=True,size=9 if i==7 else 10); c.border=BOX
            c.alignment=CEN if i in (1,10,11) else WRAP
        ws.cell(r,9).alignment=Alignment(horizontal='right',vertical='center')
        ws.row_dimensions[r].height=30
        r+=1
        if not sub:
            ws.cell(r,1,key); ws.cell(r,8,'当月該当業務なし')
            for i in range(1,13):
                c=ws.cell(r,i); c.border=BOX; c.fill=NONE; c.font=F(color='0000FF' if i==1 else '808080'); c.alignment=CEN if i<8 else WRAP
            r+=1; continue
        for j,x in enumerate(sub,1):
            vals=[key,None,j,datetime.date.fromisoformat(x['date']),f'=MID("日月火水木金土",WEEKDAY(D{r}),1)',
                  x['time'],x['proj'],x['txt'],x['aite'],x['tanto'],x['no'],x['note']]
            for i,v in enumerate(vals,1):
                c=ws.cell(r,i,v); c.font=F(color='0000FF' if i==1 else '000000',size=9 if i==2 else 10); c.border=BOX
                c.alignment=CEN if i in (1,2,3,4,5,6,10,11) else WRAP
            ws.cell(r,4).number_format='m"月"d"日"'; ws.cell(r,11).number_format='"№"0'
            r+=1
    last=r-1
    notes=['※ 出典：業務処理結果報告書（'+TAG+f' №{NO0}～№{NO1}）の記載内容を、現場技術業務委託共通仕様書の条文ごとに再整理したもの。',
     '※ 業務内容打合せ・翌日作業打合せ・業務処理結果報告書作成の日常定型業務は第5条として1日1行に集約。',
     '※ A列「条文番号」（青字）が振分け先。見出し・条文要旨は「条文マスタ」シートから参照。振分けを変更する場合はA列を書き換え、「条文別集計」に自動反映。']
    for i,t in enumerate(notes): ws.cell(r+1+i,1,t).font=F(size=9)
    for i,w in enumerate([13,12,5,8,4,8,34,48,18,13,8,14],1): ws.column_dimensions[L(i)].width=w
    ws.freeze_panes=f'A{HR+1}'; ws.print_title_rows=f'{HR}:{HR}'
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=9; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_area=f'A1:L{r+len(notes)}'
    ws.oddFooter.center.text='条文別整理表　&P / &N'; ws.page_margins.left=ws.page_margins.right=0.4

    # 条文別集計
    dates=sorted({x['date'] for x in rows})
    wa['A1']='条文別 実施件数集計表（'+TAG+'）'; wa['A1'].font=F(size=14,bold=True)
    wa['A2']='「条文別整理表」の件数をCOUNTIFSで集計（件数＝整理表の行数）。'; wa['A2'].font=F(size=9)
    hr=4; heads=['条文番号','章','見出し']+[None]*len(dates)+['合計件数','実施日数']; tc=4+len(dates)
    for j,d in enumerate(dates):
        wa.cell(hr,4+j,datetime.date.fromisoformat(d)).number_format='m/d'
        wa.cell(hr+1,4+j,f'=MID("日月火水木金土",WEEKDAY({L(4+j)}{hr}),1)')
    for i,h in enumerate(heads,1):
        if h: wa.cell(hr,i,h)
        for rr in (hr,hr+1):
            c=wa.cell(rr,i); c.font=F(bold=True,size=9); c.fill=HDR; c.border=BOX; c.alignment=CEN
    for col in (1,2,3,tc,tc+1): wa.merge_cells(start_row=hr,start_column=col,end_row=hr+1,end_column=col)
    rng=lambda col: f"条文別整理表!${col}${HR+1}:${col}${last}"
    r0=hr+2
    for n,(key,ch,title,_) in enumerate(CL):
        rr=r0+n; wa.cell(rr,1,key); wa.cell(rr,2,look('B',f'$A{rr}')); wa.cell(rr,3,look('C',f'$A{rr}'))
        for j in range(len(dates)):
            c=wa.cell(rr,4+j,f'=COUNTIFS({rng("A")},$A{rr},{rng("D")},{L(4+j)}${hr},{rng("C")},">0")'); c.number_format='0;-0;""'
        wa.cell(rr,tc,f'=SUM({L(4)}{rr}:{L(tc-1)}{rr})').number_format='0;-0;"－"'
        wa.cell(rr,tc+1,f'=COUNTIF({L(4)}{rr}:{L(tc-1)}{rr},">0")').number_format='0;-0;"－"'
        for i in range(1,tc+2):
            c=wa.cell(rr,i); c.border=BOX; c.font=F(size=9); c.alignment=CEN if i not in (2,3) else WRAP
    rt=r0+N; wa.cell(rt,1,'合計'); wa.merge_cells(start_row=rt,start_column=1,end_row=rt,end_column=3)
    for j in range(len(dates)+1): wa.cell(rt,4+j,f'=SUM({L(4+j)}{r0}:{L(4+j)}{rt-1})')
    wa.cell(rt,tc+1,f'=SUMPRODUCT(--({L(4)}{rt}:{L(tc-1)}{rt}>0))')
    for i in range(1,tc+2):
        c=wa.cell(rt,i); c.border=BOX; c.font=F(size=9,bold=True); c.fill=SEC; c.alignment=CEN
    wa.cell(rt+2,1,f'※ 合計行の「実施日数」は{M}月の業務実施日数（履行日数）。「－」は当月該当業務なし。').font=F(size=9)
    for col,w in zip('ABC',(13,20,30)): wa.column_dimensions[col].width=w
    for j in range(len(dates)): wa.column_dimensions[L(4+j)].width=4.5
    wa.column_dimensions[L(tc)].width=8; wa.column_dimensions[L(tc+1)].width=8
    wa.page_setup.orientation='landscape'; wa.page_setup.paperSize=9; wa.page_setup.fitToWidth=1; wa.page_setup.fitToHeight=1; wa.sheet_properties.pageSetUpPr.fitToPage=True


    # ---------------- 目次
    import os
    PAGES=PAGES_IN
    wt=wb.create_sheet('目次',0)
    wt['A1']='現場技術業務　検査受検資料　目次'; wt['A1'].font=F(size=16,bold=True)
    wt.merge_cells('A1:H1'); wt['A1'].alignment=Alignment(horizontal='center')
    wt['A2']=TAG; wt['A2'].font=F(size=13,bold=True); wt.merge_cells('A2:H2'); wt['A2'].alignment=Alignment(horizontal='center')
    for i,(k,v) in enumerate(info,4):
        wt.cell(i,1,k).font=F(bold=True); wt.merge_cells(start_row=i,start_column=1,end_row=i,end_column=2)
        wt.cell(i,3,v).font=F(); wt.merge_cells(start_row=i,start_column=3,end_row=i,end_column=8)
    def hdr(rr,heads):
        for i,h in enumerate(heads,1):
            c=wt.cell(rr,i,h); c.font=F(bold=True); c.fill=HDR; c.border=BOX; c.alignment=CEN
    # 1. 綴り構成
    r=9; wt.cell(r,1,'１．資料構成').font=F(size=12,bold=True); r+=1
    hdr(r,['No.','資料名','','','内容','','','備考'])
    wt.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); wt.merge_cells(start_row=r,start_column=5,end_row=r,end_column=7)
    docs=[('1','目次','本表','',None),
          ('2','条文別整理表','業務処理結果報告書の実施業務を共通仕様書の条文別に整理','',"条文別整理表"),
          ('3','条文別集計','条文別・実施日別の件数集計','',"条文別集計"),
          ('4','条文マスタ','現場技術業務委託共通仕様書 条文一覧・要旨','',"条文マスタ"),
          ('5','業務処理結果報告書',f'令和{RY}年{d0.month}月{d0.day}日～{dl.month}月{dl.day}日（№{NO0}～№{NO1}、{NDAYS}日分）','別綴り（日報原本）',None)]
    for no,nm,ct,bk,link in docs:
        r+=1
        for i,v in [(1,no),(2,nm),(5,ct),(8,bk)]: wt.cell(r,i,v)
        wt.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); wt.merge_cells(start_row=r,start_column=5,end_row=r,end_column=7)
        for i in range(1,9):
            c=wt.cell(r,i); c.border=BOX; c.font=F(); c.alignment=CEN if i==1 else WRAP
        if link:
            c=wt.cell(r,2); c.hyperlink=f"#'{link}'!A1"; c.font=F(color='0563C1',underline='single')
        wt.row_dimensions[r].height=20
    # 2. 条文別索引
    r+=2; wt.cell(r,1,'２．条文別索引（条文別整理表）').font=F(size=12,bold=True); r+=1
    hdr(r,['条文番号','見出し','件数','実施日数','実施日','該当 業務処理結果報告書№','整理表 頁','備考']); hr0=r
    for key,ch,title,_ in CL:
        r+=1; sub=[x for x in rows if x['c']==key]
        ds=sorted({x['date'] for x in sub}); nos=sorted({x['no'] for x in sub})
        if key=='第5条' and len(ds)==NDAYS: dtxt=f'全履行日（{NDAYS}日）'; ntxt=f'№{NO0}～№{NO1}'
        else:
            dtxt='、'.join(f"{int(d[5:7])}/{int(d[8:])}" for d in ds) or '―'
            ntxt='、'.join(f'№{n}' for n in nos) or '―'
        hr_=HROWS[key]; m=f'条文別集計!$A$6:$A${5+N}'
        vals=[key,look('C',f'$A{r}'),f'=INDEX(条文別集計!${L(tc)}$6:${L(tc)}${5+N},MATCH($A{r},{m},0))',
              f'=INDEX(条文別集計!${L(tc+1)}$6:${L(tc+1)}${5+N},MATCH($A{r},{m},0))',dtxt,ntxt,PAGES.get(key,''),'当月該当なし' if not sub else '']
        for i,v in enumerate(vals,1):
            c=wt.cell(r,i,v); c.border=BOX; c.font=F(color='808080' if not sub else '000000'); c.alignment=CEN if i in (1,3,4,7) else WRAP
        for i in (3,4): wt.cell(r,i).number_format='0;-0;"－"'
        wt.cell(r,7).number_format='"p."0'
        c=wt.cell(r,1); c.hyperlink=f"#'条文別整理表'!A{hr_}"; c.font=F(color='0563C1',underline='single')
        wt.row_dimensions[r].height=max(20,14*((len(title)//17)+1),14*((len(dtxt)//22)+1),14*((len(ntxt)//22)+1))
    r+=1; wt.cell(r,1,'合計'); wt.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
    wt.cell(r,3,f'=SUM(C{hr0+1}:C{r-1})'); wt.cell(r,4,f'=条文別集計!{L(tc+1)}{6+N}')
    for i in range(1,9):
        c=wt.cell(r,i); c.border=BOX; c.fill=SEC; c.font=F(bold=True); c.alignment=CEN
    notes2=['※ 条文番号をクリックすると「条文別整理表」の該当箇所へ移動。件数・実施日数は「条文別集計」から参照。',
            '※ 「整理表 頁」は条文別整理表をA4横で印刷した際のページ番号（フッター「条文別整理表 p / n」に対応）。',
            '※ 日常定型業務（業務内容打合せ・翌日作業打合せ・報告書作成）は第5条に1日1件として計上。']
    for i,t in enumerate(notes2): wt.cell(r+2+i,1,t).font=F(size=9)
    for i,w in enumerate([14,34,6,8,35,35,7,16],1): wt.column_dimensions[L(i)].width=w
    wt.page_setup.orientation='landscape'; wt.page_setup.paperSize=9; wt.page_setup.fitToWidth=1; wt.page_setup.fitToHeight=1
    wt.sheet_properties.pageSetUpPr.fitToPage=True; wt.print_area=f'A1:H{r+2+len(notes2)}'
    wb.active=0

    out=os.path.join(OUTDIR,f'現場技術業務_検査受検資料_条文別整理表_{M}月分.xlsx'); wb.calculation.fullCalcOnLoad=True; wb.save(out)
    pass
    return out, HROWS, last

def page_map(out, keys):
    if not shutil.which('soffice'): return {}
    tmp=tempfile.mkdtemp()
    wb=openpyxl.load_workbook(out)
    for n in wb.sheetnames:
        if n!='条文別整理表': del wb[n]
    p=os.path.join(tmp,'only.xlsx'); wb.save(p)
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',tmp,p],capture_output=True,timeout=300)
    pdf=os.path.join(tmp,'only.pdf')
    if not os.path.exists(pdf): return {}
    n=int(re.search(r'Pages:\s+(\d+)',subprocess.run(['pdfinfo',pdf],capture_output=True,text=True).stdout)[1])
    pages={}
    for pg in range(1,n+1):
        t=subprocess.run(['pdftotext','-layout','-f',str(pg),'-l',str(pg),pdf,'-'],capture_output=True,text=True).stdout
        for k in keys:
            if k not in pages and re.search(r'(?<![\d条])'+re.escape(k)+r'(?![第\d])',t): pages[k]=pg
    shutil.rmtree(tmp,ignore_errors=True)
    return pages

out,HROWS,_=build({})
pages=page_map(out,list(HROWS))
if pages: out,_,_=build(pages)
else: print('※ soffice/pdftotext が無いため目次の「整理表 頁」は空欄です。')
print(out)
