import sys, json, re
src=open('make_joubun.py').read(); src=src[:src.index('SRC=sys.argv')]
g={}; exec(src,g)
import openpyxl
out=[]
for f in ['a.xlsx','b.xlsx','src.xlsx']:
    ents,days,_,_=g['parse'](f)
    ws=openpyxl.load_workbook(f).active
    starts=[c.row for c in ws['H'] if c.value=='業務処理結果報告書']
    wx={no:str(ws.cell(s+19,23).value or '').strip() for s,(d,no) in zip(starts,days)}
    for e in ents:
        e['wx']=wx[e['no']]
        if e['cat'] in ('業務内容打合せ','業務打合せ','『業務処理結果報告書』作成') or e['cat'].startswith('執務室') or e['cat'] in ('河川砂防課','書面資料打合せ'): continue
        out.append(e)
json.dump(out,open('field.json','w'),ensure_ascii=False)
for e in out: print(e['no'],e['date'][5:],e['wx'],e['time'],'|',e['cat'],'|',e['det'],'|',e['aite'],'|',e['tanto'])
